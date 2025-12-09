# -*- coding: utf-8 -*-

import io
import base64
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from odoo import fields, models, api, _


def create_attachment(file_type):
    if file_type == 'docx':
        doc = Document()
        output = io.BytesIO()
        doc.save(output)
    elif file_type == 'xlsx':
        wb = Workbook()
        output = io.BytesIO()
        wb.save(output)
    elif file_type == 'pptx':
        prs = Presentation()
        output = io.BytesIO()
        prs.save(output)
    else:
        return False

    file_data = output.getvalue()
    file_base64 = base64.b64encode(file_data)
    return file_base64


class DocumentsCreate(models.TransientModel):
    _name = 'documents.create'
    _description = 'Documents Create'

    name = fields.Char(string='Name')
    type = fields.Selection(string='Type',
                            selection=[('docx', '📘 Docs'),
                                       ('xlsx', '📗 Excel'),
                                       ('pptx', '📙 PowerPoint'),
                                       ], default='docx')
    document_directory_id = fields.Many2one('document.directory', string="Directory", ondelete="cascade")

    def action_document_create(self):
        self.env['ir.attachment'].create({
            'name': f'{self.name}.{self.type}',
            'type': 'binary',
            'datas': create_attachment(self.type),
            'document_directory_id': self.document_directory_id.id,
        })

        return {
            'name': _(self.document_directory_id.name),
            'type': 'ir.actions.act_window',
            'view_mode': 'kanban,tree,form',
            'res_model': 'ir.attachment',
            'target': 'current',
            'domain': [('document_directory_id', '=', self.document_directory_id.id)],
            'context': {'default_document_directory_id': self.document_directory_id.id},
        }
